"""Generate the Personnel sheet.

Section 1 (top): Personnel table with info + auto-calculated leave balances
Section 2 (bottom): Leave Credits / Transactions log
"""

from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName


# --- Styles ---
TITLE_FONT = Font(bold=True, size=16)
SECTION_FONT = Font(bold=True, size=13, color="FFFFFF")
SECTION_FILL = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
HEADER_FONT = Font(bold=True, size=10, color="FFFFFF")
HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
VL_HDR = PatternFill(start_color="548235", end_color="548235", fill_type="solid")
PHOL_HDR = PatternFill(start_color="BF8F00", end_color="BF8F00", fill_type="solid")
OIL_HDR = PatternFill(start_color="C00000", end_color="C00000", fill_type="solid")
VL_FILL = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
PHOL_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
OIL_FILL = PatternFill(start_color="FCE4EC", end_color="FCE4EC", fill_type="solid")
TOTAL_FILL = PatternFill(start_color="D9E2F3", end_color="D9E2F3", fill_type="solid")
THIN_BORDER = Border(
    left=Side(style="thin", color="B4B4B4"),
    right=Side(style="thin", color="B4B4B4"),
    top=Side(style="thin", color="B4B4B4"),
    bottom=Side(style="thin", color="B4B4B4"),
)

# Personnel table columns
P_HEADERS = [
    "Name", "Rank", "Shift", "Join Date", "ORD Date", "Status",
    "VL Ent", "VL Used", "VL Bal",
    "PHOL Cr", "PHOL Used", "PHOL Bal",
    "OIL Cr", "OIL Used", "OIL Bal",
    "Total Bal",
]
P_HEADER_FILLS = [
    HEADER_FILL, HEADER_FILL, HEADER_FILL, HEADER_FILL, HEADER_FILL, HEADER_FILL,
    VL_HDR, VL_HDR, VL_HDR,
    PHOL_HDR, PHOL_HDR, PHOL_HDR,
    OIL_HDR, OIL_HDR, OIL_HDR,
    HEADER_FILL,
]
P_COL_WIDTHS = [20, 6, 6, 12, 12, 8, 7, 8, 7, 8, 9, 8, 7, 8, 7, 9]

# Leave Credits table columns
LC_HEADERS = [
    "#", "Name", "Txn", "Leave", "Days", "Date", "Expiry", "Source#", "Cal Ref", "Remarks",
]
LC_COL_WIDTHS = [5, 18, 8, 7, 6, 12, 12, 8, 12, 30]

# The row where the Personnel data starts (after title + headers)
PERSONNEL_DATA_START = 4
PERSONNEL_MAX_ROWS = 50  # max 50 personnel rows
# The row where Leave Credits section starts
CREDITS_SECTION_START = PERSONNEL_DATA_START + PERSONNEL_MAX_ROWS + 3


def create_personnel_sheet(wb):
    ws = wb.create_sheet("Personnel")

    # =============================================
    # SECTION 1: Personnel Info
    # =============================================
    ws.cell(row=1, column=1, value="Personnel").font = TITLE_FONT
    ws.merge_cells("A1:P1")

    # Sub-header grouping labels
    ws.cell(row=2, column=1, value="Information").font = Font(bold=True, size=10)
    ws.merge_cells("A2:F2")
    ws.cell(row=2, column=7, value="Vacation Leave").font = Font(bold=True, size=10, color="548235")
    ws.merge_cells("G2:I2")
    ws.cell(row=2, column=10, value="PH Off-in-Lieu").font = Font(bold=True, size=10, color="BF8F00")
    ws.merge_cells("J2:L2")
    ws.cell(row=2, column=13, value="Off-in-Lieu").font = Font(bold=True, size=10, color="C00000")
    ws.merge_cells("M2:O2")

    # Column headers (row 3)
    for col_idx, (header, hfill) in enumerate(zip(P_HEADERS, P_HEADER_FILLS), 1):
        cell = ws.cell(row=3, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = hfill
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        cell.border = THIN_BORDER

    # Pre-format data rows with borders and light background
    for row in range(PERSONNEL_DATA_START, PERSONNEL_DATA_START + PERSONNEL_MAX_ROWS):
        for col in range(1, 17):
            ws.cell(row=row, column=col).border = THIN_BORDER
        # Date format
        ws.cell(row=row, column=4).number_format = "DD-MMM-YYYY"
        ws.cell(row=row, column=5).number_format = "DD-MMM-YYYY"
        # Light background for leave columns
        for c in [7, 8, 9]:
            ws.cell(row=row, column=c).fill = VL_FILL
        for c in [10, 11, 12]:
            ws.cell(row=row, column=c).fill = PHOL_FILL
        for c in [13, 14, 15]:
            ws.cell(row=row, column=c).fill = OIL_FILL
        ws.cell(row=row, column=16).fill = TOTAL_FILL

    # Data validation for Shift
    dv_shift = DataValidation(type="list", formula1='"1,2,3"', allow_blank=True)
    ws.add_data_validation(dv_shift)
    dv_shift.add(f"C{PERSONNEL_DATA_START}:C{PERSONNEL_DATA_START + PERSONNEL_MAX_ROWS - 1}")

    # Data validation for Status
    dv_status = DataValidation(type="list", formula1='"Active,ORD,Posted"', allow_blank=True)
    ws.add_data_validation(dv_status)
    dv_status.add(f"F{PERSONNEL_DATA_START}:F{PERSONNEL_DATA_START + PERSONNEL_MAX_ROWS - 1}")

    # Column widths
    for i, w in enumerate(P_COL_WIDTHS):
        ws.column_dimensions[chr(65 + i)].width = w

    # Note: Balance columns (H, I, K, L, N, O, P) are populated by VBA.
    # VBA reads from the Leave Credits table below and calculates
    # expiry-aware balances per person.
    # A note for the user:
    note_row = PERSONNEL_DATA_START + PERSONNEL_MAX_ROWS + 1
    ws.cell(row=note_row, column=1,
            value="Balances (Used/Bal columns) are auto-calculated by VBA. Click Refresh or edit the Calendar to update."
    ).font = Font(italic=True, color="888888", size=9)
    ws.merge_cells(f"A{note_row}:P{note_row}")

    # =============================================
    # SECTION 2: Leave Credits / Transactions
    # =============================================
    sec2 = CREDITS_SECTION_START

    ws.cell(row=sec2, column=1, value="Leave Credits & Transactions").font = SECTION_FONT
    ws.cell(row=sec2, column=1).fill = SECTION_FILL
    for c in range(1, 11):
        ws.cell(row=sec2, column=c).fill = SECTION_FILL
    ws.merge_cells(f"A{sec2}:J{sec2}")
    ws.row_dimensions[sec2].height = 22

    sec2 += 1
    ws.cell(row=sec2, column=1,
            value="Add PHOL/OIL credits below. VL is auto-credited when personnel are added. "
                  "DEBIT rows are auto-created when leave is plotted on the Calendar."
    ).font = Font(italic=True, color="666666", size=9)
    ws.merge_cells(f"A{sec2}:J{sec2}")

    # Column headers
    hdr_row = sec2 + 1
    for col_idx, header in enumerate(LC_HEADERS, 1):
        cell = ws.cell(row=hdr_row, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        cell.border = THIN_BORDER

    # Pre-format 200 credit rows
    data_start = hdr_row + 1
    for row in range(data_start, data_start + 200):
        for col in range(1, 11):
            ws.cell(row=row, column=col).border = THIN_BORDER
        ws.cell(row=row, column=6).number_format = "DD-MMM-YYYY"
        ws.cell(row=row, column=7).number_format = "DD-MMM-YYYY"
        ws.cell(row=row, column=9).number_format = "DD-MMM-YYYY"

    # Data validation for Txn type
    dv_txn = DataValidation(type="list", formula1='"CREDIT,DEBIT"', allow_blank=True)
    ws.add_data_validation(dv_txn)
    dv_txn.add(f"C{data_start}:C{data_start + 199}")

    # Data validation for Leave type
    dv_ltype = DataValidation(type="list", formula1='"VL,PHOL,OIL"', allow_blank=True)
    ws.add_data_validation(dv_ltype)
    dv_ltype.add(f"D{data_start}:D{data_start + 199}")

    # Credit column widths
    for i, w in enumerate(LC_COL_WIDTHS):
        ws.column_dimensions[chr(65 + i)].width = max(
            ws.column_dimensions[chr(65 + i)].width or 0, w
        )

    # Store key row references as named ranges for VBA
    def _def(name, ref):
        wb.defined_names.add(DefinedName(name, attr_text=ref))

    _def("PersonnelDataStart", f"Personnel!$A${PERSONNEL_DATA_START}")
    _def("CreditsHeaderRow", f"Personnel!$A${hdr_row}")
    _def("CreditsDataStart", f"Personnel!$A${data_start}")

    # Freeze at the data area
    ws.freeze_panes = f"A{PERSONNEL_DATA_START}"

    return ws
