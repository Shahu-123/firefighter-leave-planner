"""Generate the Config sheet with station parameters and public holidays."""

from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.workbook.defined_name import DefinedName
from config import STATION_CONFIG, SG_PUBLIC_HOLIDAYS


def _def_name(wb, name, attr_text):
    wb.defined_names.add(DefinedName(name, attr_text=attr_text))


HEADER_FONT = Font(bold=True, size=11)
HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
HEADER_FONT_WHITE = Font(bold=True, size=11, color="FFFFFF")
THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)


def create_config_sheet(wb):
    ws = wb.create_sheet("Config")

    # --- Section: Station Parameters ---
    ws["A1"] = "Station Parameters"
    ws["A1"].font = Font(bold=True, size=14)
    ws.merge_cells("A1:B1")

    params = [
        ("StationName", STATION_CONFIG["station_name"]),
        ("Year", STATION_CONFIG["year"]),
        ("CycleAnchorDate", STATION_CONFIG["anchor_date"]),
        ("AnchorShift", STATION_CONFIG["anchor_shift"]),
        ("SlotsL1", STATION_CONFIG["slots_l1"]),
        ("SlotsL2", STATION_CONFIG["slots_l2"]),
        ("SlotsKIV", STATION_CONFIG["slots_kiv"]),
        ("VLDaysFullYear", STATION_CONFIG["vl_days_full_year"]),
        ("OILExpiryMonths", STATION_CONFIG["oil_expiry_months"]),
        ("PHOLExpiryMonths", STATION_CONFIG["phol_expiry_months"]),
        ("PlotWindowMonths", STATION_CONFIG["plot_window_months"]),
        ("ShiftLeaveCostDays", STATION_CONFIG["shift_leave_cost_days"]),
    ]

    # Headers
    for col, header in enumerate(["Parameter", "Value"], 1):
        cell = ws.cell(row=2, column=col, value=header)
        cell.font = HEADER_FONT_WHITE
        cell.fill = HEADER_FILL
        cell.border = THIN_BORDER

    for i, (key, val) in enumerate(params, start=3):
        ws.cell(row=i, column=1, value=key).border = THIN_BORDER
        ws.cell(row=i, column=1).font = Font(bold=True)
        cell = ws.cell(row=i, column=2, value=val)
        cell.border = THIN_BORDER
        if isinstance(val, int):
            cell.number_format = "0"
        # Create named range for each parameter
        _def_name(wb,f"cfg_{key}", attr_text=f"Config!$B${i}")

    # --- Section: Public Holidays ---
    ph_start_row = len(params) + 5
    ws.cell(row=ph_start_row, column=1, value="Public Holidays 2026")
    ws.cell(row=ph_start_row, column=1).font = Font(bold=True, size=14)
    ws.merge_cells(f"A{ph_start_row}:B{ph_start_row}")

    ph_header_row = ph_start_row + 1
    for col, header in enumerate(["Date", "Holiday Name"], 1):
        cell = ws.cell(row=ph_header_row, column=col, value=header)
        cell.font = HEADER_FONT_WHITE
        cell.fill = HEADER_FILL
        cell.border = THIN_BORDER

    for i, (ph_date, ph_name) in enumerate(SG_PUBLIC_HOLIDAYS, start=ph_header_row + 1):
        cell_d = ws.cell(row=i, column=1, value=ph_date)
        cell_d.number_format = "DD-MMM-YYYY"
        cell_d.border = THIN_BORDER
        cell_n = ws.cell(row=i, column=2, value=ph_name)
        cell_n.border = THIN_BORDER

    # Define named range for PH table
    ph_end_row = ph_header_row + len(SG_PUBLIC_HOLIDAYS)
    _def_name(wb,
        "tbl_PublicHolidays",
        attr_text=f"Config!$A${ph_header_row + 1}:$B${ph_end_row}",
    )

    # Column widths
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 30

    return ws
