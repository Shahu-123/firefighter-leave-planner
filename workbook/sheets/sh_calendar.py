"""Generate the Calendar sheet - primary leave plotting interface.

Layout: one row per date (Apr-Dec 2026) with month separator rows.
Slot columns (D-I) have data validation dropdowns filtered by on-duty shift.
Users plot leave by selecting a name from the dropdown.
"""

import calendar as cal_mod
from datetime import date, timedelta
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName
from config import STATION_CONFIG, SG_PUBLIC_HOLIDAYS, get_shift_for_date


# --- Styles ---
TITLE_FONT = Font(bold=True, size=16)
MONTH_FONT = Font(bold=True, size=13, color="FFFFFF")
MONTH_FILL = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
HEADER_FONT = Font(bold=True, size=10, color="FFFFFF")
HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")

L1_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
L2_FILL = PatternFill(start_color="FFE699", end_color="FFE699", fill_type="solid")
KIV_FILL = PatternFill(start_color="F4CCCC", end_color="F4CCCC", fill_type="solid")
PH_FILL = PatternFill(start_color="DAEEF3", end_color="DAEEF3", fill_type="solid")
SAT_FILL = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
SUN_FILL = PatternFill(start_color="E8E8E8", end_color="E8E8E8", fill_type="solid")
L1_HDR = PatternFill(start_color="548235", end_color="548235", fill_type="solid")
L2_HDR = PatternFill(start_color="BF8F00", end_color="BF8F00", fill_type="solid")
KIV_HDR = PatternFill(start_color="C00000", end_color="C00000", fill_type="solid")

THIN_BORDER = Border(
    left=Side(style="thin", color="B4B4B4"),
    right=Side(style="thin", color="B4B4B4"),
    top=Side(style="thin", color="B4B4B4"),
    bottom=Side(style="thin", color="B4B4B4"),
)

MONTH_NAMES = [
    "", "January", "February", "March", "April", "May", "June",
    "July", "August", "September", "October", "November", "December",
]

HEADERS = ["Date", "Day", "Shift", "L1", "L1", "L2", "L2", "KIV", "KIV", "On Lv", "Avail", "Remarks"]
HEADER_FILLS = [
    HEADER_FILL, HEADER_FILL, HEADER_FILL,
    L1_HDR, L1_HDR, L2_HDR, L2_HDR, KIV_HDR, KIV_HDR,
    HEADER_FILL, HEADER_FILL, HEADER_FILL,
]
COL_WIDTHS = [10, 5, 5, 18, 18, 18, 18, 18, 18, 6, 6, 18]


def create_calendar_sheet(wb):
    ws = wb.create_sheet("Calendar")

    year = STATION_CONFIG["year"]
    start_month = STATION_CONFIG["start_month"]
    ph_lookup = {d: name for d, name in SG_PUBLIC_HOLIDAYS}
    day_abbr = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]

    # --- Data validation: dropdown linked to shift name lists ---
    # These named ranges (Shift1_Names etc.) are populated by VBA on open.
    # We create 3 separate DV objects, one per shift.
    dv_map = {}
    for s in [1, 2, 3]:
        dv = DataValidation(
            type="list",
            formula1=f"=Shift{s}_Names",
            allow_blank=True,
            showErrorMessage=True,
            errorTitle="Invalid Name",
            error=f"Select a name from Shift {s}, or leave blank.",
        )
        dv.showDropDown = False  # False = show dropdown arrow (counterintuitive API)
        ws.add_data_validation(dv)
        dv_map[s] = dv

    # --- Column widths ---
    for i, w in enumerate(COL_WIDTHS):
        ws.column_dimensions[chr(65 + i)].width = w

    row = 1

    # --- Title ---
    ws.cell(row=row, column=1, value=f"STATION Leave Calendar {year}").font = TITLE_FONT
    ws.merge_cells(f"A{row}:L{row}")
    row += 1

    # --- Build calendar month by month ---
    # Track which rows are data rows (for DV) vs separator rows
    # We store a mapping: row -> shift (for applying correct DV)
    row_shift_map = {}

    for month in range(start_month, 13):
        # Month separator
        row += 1
        ws.cell(row=row, column=1, value=f"{MONTH_NAMES[month]} {year}").font = MONTH_FONT
        ws.cell(row=row, column=1).fill = MONTH_FILL
        ws.cell(row=row, column=1).alignment = Alignment(vertical="center")
        for c in range(1, 13):
            ws.cell(row=row, column=c).fill = MONTH_FILL
        ws.merge_cells(f"A{row}:L{row}")
        ws.row_dimensions[row].height = 22
        row += 1

        # Column headers (repeated per month for readability)
        for col_idx, (header, hfill) in enumerate(zip(HEADERS, HEADER_FILLS), 1):
            cell = ws.cell(row=row, column=col_idx, value=header)
            cell.font = HEADER_FONT
            cell.fill = hfill
            cell.alignment = Alignment(horizontal="center")
            cell.border = THIN_BORDER
        row += 1

        # Date rows
        days_in_month = cal_mod.monthrange(year, month)[1]
        for day in range(1, days_in_month + 1):
            d = date(year, month, day)
            shift = get_shift_for_date(d)
            day_name = day_abbr[d.weekday()]
            is_ph = d in ph_lookup

            ws.cell(row=row, column=1, value=d).number_format = "DD-MMM"
            ws.cell(row=row, column=2, value=day_name)
            ws.cell(row=row, column=3, value=shift)

            # Slot cells D-I: apply slot background colors
            for c in [4, 5]:
                ws.cell(row=row, column=c).fill = L1_FILL
            for c in [6, 7]:
                ws.cell(row=row, column=c).fill = L2_FILL
            for c in [8, 9]:
                ws.cell(row=row, column=c).fill = KIV_FILL

            # On Leave count (col J)
            ws.cell(row=row, column=10).value = f"=COUNTA(D{row}:I{row})"
            ws.cell(row=row, column=10).alignment = Alignment(horizontal="center")

            # Available Strength (col K)
            ws.cell(row=row, column=11).value = (
                f'=COUNTIFS(Personnel!C$4:C$100,C{row},'
                f'Personnel!F$4:F$100,"Active")-J{row}'
            )
            ws.cell(row=row, column=11).alignment = Alignment(horizontal="center")

            # Remarks (col L) - PH name or manual notes
            if is_ph:
                ws.cell(row=row, column=12, value=ph_lookup[d])

            # Borders
            for c in range(1, 13):
                ws.cell(row=row, column=c).border = THIN_BORDER

            # Weekend/PH highlighting on date columns
            if is_ph:
                for c in [1, 2, 3, 12]:
                    ws.cell(row=row, column=c).fill = PH_FILL
            elif d.weekday() == 5:  # Saturday
                for c in [1, 2, 3]:
                    ws.cell(row=row, column=c).fill = SAT_FILL
            elif d.weekday() == 6:  # Sunday
                for c in [1, 2, 3]:
                    ws.cell(row=row, column=c).fill = SUN_FILL

            # Apply data validation for this shift to slot cells
            dv = dv_map[shift]
            for c_letter in ["D", "E", "F", "G", "H", "I"]:
                dv.add(f"{c_letter}{row}")

            row_shift_map[row] = shift
            row += 1

    # Freeze panes: freeze columns A-C and allow scrolling of slot columns
    ws.freeze_panes = "D3"

    # Legend at bottom
    row += 2
    ws.cell(row=row, column=1, value="Legend:").font = Font(bold=True, size=10)
    ws.cell(row=row, column=2, value="L1").fill = L1_FILL
    ws.cell(row=row, column=2).border = THIN_BORDER
    ws.cell(row=row, column=3, value="L2").fill = L2_FILL
    ws.cell(row=row, column=3).border = THIN_BORDER
    ws.cell(row=row, column=4, value="KIV").fill = KIV_FILL
    ws.cell(row=row, column=4).border = THIN_BORDER
    ws.cell(row=row, column=5, value="PH").fill = PH_FILL
    ws.cell(row=row, column=5).border = THIN_BORDER
    row += 1
    ws.cell(row=row, column=1, value="Select names from dropdown. VBA handles leave deduction & slot bumping.").font = Font(
        italic=True, color="666666", size=9
    )

    return ws
