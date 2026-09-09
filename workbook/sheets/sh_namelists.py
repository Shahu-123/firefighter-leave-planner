"""Generate the hidden NameLists sheet for dropdown data validation.

Contains 3 columns: Shift 1 names, Shift 2 names, Shift 3 names.
Populated by VBA on workbook open and after personnel changes.
Named ranges Shift1_Names, Shift2_Names, Shift3_Names point here.
"""

from openpyxl.styles import Font
from openpyxl.workbook.defined_name import DefinedName


MAX_NAMES = 50  # generous max per shift


def create_namelists_sheet(wb):
    ws = wb.create_sheet("NameLists")

    # Headers
    ws.cell(row=1, column=1, value="Shift 1").font = Font(bold=True)
    ws.cell(row=1, column=2, value="Shift 2").font = Font(bold=True)
    ws.cell(row=1, column=3, value="Shift 3").font = Font(bold=True)

    # Placeholder data (VBA will overwrite on open)
    ws.cell(row=2, column=1, value="(Add personnel first)")
    ws.cell(row=2, column=2, value="(Add personnel first)")
    ws.cell(row=2, column=3, value="(Add personnel first)")

    # Define named ranges for each shift
    # These are initially set to the full range; VBA will trim to actual data
    for s, col_letter in [(1, "A"), (2, "B"), (3, "C")]:
        dn = DefinedName(
            f"Shift{s}_Names",
            attr_text=f"NameLists!${col_letter}$2:${col_letter}${MAX_NAMES + 1}",
        )
        wb.defined_names.add(dn)

    # Hide the sheet
    ws.sheet_state = "hidden"

    return ws
